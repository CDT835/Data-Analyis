"""Xuất báo cáo (danh sách khối) ra Word, PDF và Excel.

Word: python-docx, Times New Roman, chữ đen nền trắng, không có khung màu trang trí.
PDF: reportlab, phông DejaVu Serif lấy từ matplotlib để hiển thị đúng tiếng Việt.
Excel: openpyxl, có biểu đồ Excel thật; số được ghi dạng số để người dùng tự tính tiếp.
Biểu đồ trong Word và PDF vẽ bằng matplotlib (nền Agg).
"""
import io
import os
import re
from typing import Any, Dict, List
from xml.sax.saxutils import escape

PALETTE = ["#1f3a5f", "#8a1c1c", "#2e6b3a", "#6b4e16", "#555555", "#4b2a6b", "#0f6b6b", "#999999"]
CALLOUT_LABEL = {"warn": "Cảnh báo", "info": "Lưu ý", "ok": "Kết quả"}
MIME = {"docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "pdf": "application/pdf",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}


def safe_filename(title: str, ext: str) -> str:
    base = re.sub(r"[^\w\- ]", "", title, flags=re.UNICODE).strip().replace(" ", "_")[:80] or "bao_cao"
    return f"{base}.{ext}"


# ====================================================================== biểu đồ (matplotlib)
def chart_png(b: Dict[str, Any], dpi: int = 150) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    labels, series = b["labels"], b["series"]
    fig, ax = plt.subplots(figsize=(7.2, 3.6))
    try:
        x = list(range(len(labels)))
        if b["kind"] == "bar":
            w = 0.8 / max(1, len(series))
            for i, s in enumerate(series):
                ax.bar([xi + i * w - 0.4 + w / 2 for xi in x], [v if v is not None else 0 for v in s["values"]], w, label=s["name"], color=PALETTE[i % len(PALETTE)])
        elif b["kind"] == "stacked":
            bottom = [0.0] * len(labels)
            for i, s in enumerate(series):
                vals = [v or 0.0 for v in s["values"]]
                ax.bar(x, vals, 0.6, bottom=bottom, label=s["name"], color=PALETTE[i % len(PALETTE)])
                bottom = [a + c for a, c in zip(bottom, vals)]
        else:
            for i, s in enumerate(series):
                ax.plot(x, [v if v is not None else float("nan") for v in s["values"]], marker="o", linewidth=1.8, label=s["name"], color=PALETTE[i % len(PALETTE)])
        ax.set_xticks(x)
        ax.set_xticklabels(labels, fontsize=8)
        ax.tick_params(axis="y", labelsize=8)
        if b.get("unit"):
            ax.set_ylabel(b["unit"], fontsize=8)
        ax.grid(axis="y", alpha=0.3)
        ax.spines[["top", "right"]].set_visible(False)
        ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _p: f"{v:,.0f}".replace(",", ".") if abs(v) >= 100 else f"{v:g}".replace(".", ",")))
        ax.legend(fontsize=7, loc="upper center", bbox_to_anchor=(0.5, -0.12), ncol=min(3, len(series)), frameon=False)
        fig.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=dpi)
        return buf.getvalue()
    finally:
        plt.close(fig)


# ====================================================================== Word
def export_docx(rep: Dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    usable = 17.0

    def font(style, size, bold=None):
        f = style.font
        f.name = "Times New Roman"
        f.size = Pt(size)
        f.color.rgb = RGBColor(0, 0, 0)
        if bold is not None:
            f.bold = bold
        rpr = style.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]  # bỏ phông theme (Calibri/Cambria) để tiêu đề cũng dùng Times New Roman
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), "Times New Roman")

    font(doc.styles["Normal"], 14)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    font(doc.styles["Heading 1"], 18, True)
    font(doc.styles["Heading 2"], 16, True)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    font(doc.styles["List Bullet"], 14)

    def run(p, text, size=None, bold=None, italic=None):
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        if size:
            r.font.size = Pt(size)
        if bold is not None:
            r.bold = bold
        if italic is not None:
            r.italic = italic
        r.font.color.rgb = RGBColor(0, 0, 0)
        return r

    # số trang ở chân trang
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r = fp.add_run()
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        r._r.append(el)

    def set_cell(cell, text, size, bold=False, align=None):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if align:
            p.alignment = align
        run(p, text, size, bold)

    def table(b):
        headers, rows = b["headers"], b["rows"]
        n = len(headers)
        size = 12 if n <= 5 else (11 if n <= 7 else 10)
        if b.get("caption"):
            cp = doc.add_paragraph()
            cp.paragraph_format.keep_with_next = True
            run(cp, b["caption"], 13, True)
        t = doc.add_table(rows=1, cols=n)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if n == 1:
            widths = [usable]
        elif n == 2:
            widths = [usable * 0.35, usable * 0.65]
        elif n == 3:
            widths = [usable * 0.44, usable * 0.28, usable * 0.28]
        else:
            first = 6.0 if n <= 6 else 5.0  # cột đầu rộng để tên chỉ số không bị xuống dòng nhiều
            widths = [first] + [(usable - first) / (n - 1)] * (n - 1)
        num_cols = set(b["num_cols"]) if b.get("num_cols") is not None else set([] if n <= 2 and not b.get("raw") else range(1, n))
        if n == 2 and not b.get("raw"):
            num_cols = set()
        for j, h in enumerate(headers):
            set_cell(t.rows[0].cells[j], h, size, True, WD_ALIGN_PARAGRAPH.CENTER)
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        for row in rows:
            cells = t.add_row().cells
            for j, v in enumerate(row):
                set_cell(cells[j], str(v), size, False, WD_ALIGN_PARAGRAPH.RIGHT if j in num_cols else None)
        for j, col in enumerate(t.columns):  # đặt cả độ rộng cột (lưới bảng) lẫn từng ô để Word và LibreOffice đều theo
            col.width = Cm(widths[j])
        for row in t.rows:
            for j, cell in enumerate(row.cells):
                cell.width = Cm(widths[j])
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    for b in rep["blocks"]:
        t = b["t"]
        if t == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(60)
            run(p, b["text"].upper(), 18, True)
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(p2, b.get("subtitle", ""), 14, None, True)
        elif t == "heading":
            doc.add_heading(f"{b['num']}. {b['text']}", level=b["level"])
        elif t == "para":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run(p, b["text"])
        elif t == "bullets":
            for it in b["items"]:
                p = doc.add_paragraph(style="List Bullet")
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                run(p, it)
        elif t == "table":
            table(b)
        elif t == "kpis":
            table({"headers": ["Chỉ tiêu", "Giá trị", "So sánh"], "rows": [[k["label"], k["value"], k["delta"]] for k in b["items"]], "num_cols": [1]})
        elif t == "callout":
            p = doc.add_paragraph()
            run(p, CALLOUT_LABEL.get(b["kind"], "Lưu ý") + ": ", None, True, True)
            run(p, b["text"], None, None, True)
        elif t == "chart":
            doc.add_picture(io.BytesIO(chart_png(b)), width=Cm(15.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(cp, b["fig"], 12, None, True)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ====================================================================== PDF
def export_pdf(rep: Dict[str, Any]) -> bytes:
    import matplotlib
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    fdir = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
    for name, fn in (("Serif", "DejaVuSerif.ttf"), ("Serif-B", "DejaVuSerif-Bold.ttf"), ("Serif-I", "DejaVuSerif-Italic.ttf"), ("Serif-BI", "DejaVuSerif-BoldItalic.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, os.path.join(fdir, fn)))
    pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-B", italic="Serif-I", boldItalic="Serif-BI")

    st = {
        "title": ParagraphStyle("t", fontName="Serif-B", fontSize=18, leading=24, alignment=TA_CENTER, spaceBefore=120, spaceAfter=10),
        "sub": ParagraphStyle("s", fontName="Serif-I", fontSize=12, leading=16, alignment=TA_CENTER, spaceAfter=24),
        "h1": ParagraphStyle("h1", fontName="Serif-B", fontSize=15, leading=20, spaceBefore=16, spaceAfter=8, keepWithNext=1),
        "h2": ParagraphStyle("h2", fontName="Serif-B", fontSize=13, leading=17, spaceBefore=10, spaceAfter=6, keepWithNext=1),
        "p": ParagraphStyle("p", fontName="Serif", fontSize=11, leading=15.5, alignment=TA_JUSTIFY, spaceAfter=6),
        "li": ParagraphStyle("li", fontName="Serif", fontSize=11, leading=15.5, leftIndent=14, bulletIndent=2, alignment=TA_JUSTIFY, spaceAfter=3),
        "cap": ParagraphStyle("c", fontName="Serif-B", fontSize=10.5, leading=14, spaceBefore=6, spaceAfter=4, keepWithNext=1),
        "fig": ParagraphStyle("f", fontName="Serif-I", fontSize=10, leading=13, alignment=TA_CENTER, spaceAfter=8),
        "call": ParagraphStyle("k", fontName="Serif-I", fontSize=10.5, leading=14.5, spaceAfter=6, leftIndent=8),
    }
    W = A4[0] - 4 * cm

    def tcell(text, bold=False, right=False, size=9):
        return Paragraph(escape(str(text)), ParagraphStyle("tc", fontName="Serif-B" if bold else "Serif", fontSize=size, leading=size + 3, alignment=TA_RIGHT if right else 0))

    def table(b):
        headers, rows = b["headers"], b["rows"]
        n = len(headers)
        size = 9 if n <= 6 else 8
        num_cols = set(b["num_cols"]) if b.get("num_cols") is not None else set([] if (n <= 2 and not b.get("raw")) else range(1, n))
        data = [[tcell(h, True, False, size) for h in headers]] + [[tcell(v, False, j in num_cols, size) for j, v in enumerate(r)] for r in rows]
        if n == 2:
            widths = [W * 0.35, W * 0.65]
        else:
            first = min(W * 0.42, max(W * 0.30, 4.6 * cm))
            widths = [first] + [(W - first) / (n - 1)] * (n - 1)
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                               ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        out = []
        if b.get("caption"):
            out.append(Paragraph(escape(b["caption"]), st["cap"]))
        out += [t, Spacer(1, 8)]
        return out

    flow: List[Any] = []
    for b in rep["blocks"]:
        t = b["t"]
        if t == "title":
            flow += [Paragraph(escape(b["text"].upper()), st["title"]), Paragraph(escape(b.get("subtitle", "")), st["sub"])]
        elif t == "heading":
            flow.append(Paragraph(escape(f"{b['num']}. {b['text']}"), st["h1"] if b["level"] == 1 else st["h2"]))
        elif t == "para":
            flow.append(Paragraph(escape(b["text"]), st["p"]))
        elif t == "bullets":
            flow += [Paragraph(escape(i), st["li"], bulletText="•") for i in b["items"]]
        elif t == "table":
            flow += table(b)
        elif t == "kpis":
            flow += table({"headers": ["Chỉ tiêu", "Giá trị", "So sánh"], "rows": [[k["label"], k["value"], k["delta"]] for k in b["items"]], "num_cols": [1]})
        elif t == "callout":
            flow.append(Paragraph(f"<b>{CALLOUT_LABEL.get(b['kind'], 'Lưu ý')}:</b> {escape(b['text'])}", st["call"]))
        elif t == "chart":
            img = Image(io.BytesIO(chart_png(b, 130)), width=W * 0.95, height=W * 0.95 * 0.5)
            flow.append(KeepTogether([img, Paragraph(escape(b["fig"]), st["fig"])]))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Serif", 9)
        canvas.drawCentredString(A4[0] / 2, 1.1 * cm, str(canvas.getPageNumber()))
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm, title=rep["title"], author="AI Report Studio").build(flow, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r", "\n")


def _cell(ws, row: int, col: int, value=None):
    """Ghi một ô Excel; chuỗi bắt đầu bằng = + - @ bị ép thành VĂN BẢN (chống tiêm công thức / CSV-injection)."""
    c = ws.cell(row, col, value)
    if isinstance(value, str) and value[:1] in _FORMULA_PREFIX:
        c.data_type = "s"
        c.quotePrefix = True
    return c


# ====================================================================== Excel
def export_xlsx(rep: Dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Báo cáo"
    F = lambda **k: Font(name="Times New Roman", size=k.pop("size", 12), color="000000", **k)  # noqa: E731
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ncols = 12
    ws.page_setup.orientation = "landscape"  # in ra vừa chiều ngang trang, không bị cắt cột
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.column_dimensions["A"].width = 46
    for col in "BCDEFGHIJKL":
        ws.column_dimensions[col].width = 17
    r = 1

    def text_row(text, font, height_chars=120):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        c = _cell(ws, r, 1, text)
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = max(18, 17 * (1 + len(text) // height_chars))
        r += 1

    for b in rep["blocks"]:
        t = b["t"]
        if t == "title":
            text_row(b["text"].upper(), F(size=16, bold=True))
            text_row(b.get("subtitle", ""), F(italic=True))
            r += 1
        elif t == "heading":
            r += 1
            text_row(f"{b['num']}. {b['text']}", F(size=14 if b["level"] == 1 else 13, bold=True), 90)
        elif t == "para":
            text_row(b["text"], F())
        elif t == "bullets":
            for it in b["items"]:
                text_row("• " + it, F())
        elif t == "callout":
            text_row(f"{CALLOUT_LABEL.get(b['kind'], 'Lưu ý')}: {b['text']}", F(italic=True))
        elif t in ("table", "kpis"):
            if t == "kpis":
                b = {"headers": ["Chỉ tiêu", "Giá trị", "So sánh"], "rows": [[k["label"], k["value"], k["delta"]] for k in b["items"]]}
            if b.get("caption"):
                text_row(b["caption"], F(bold=True))
            for j, h in enumerate(b["headers"][:ncols]):
                c = _cell(ws, r, j + 1, h)
                c.font, c.border, c.alignment = F(bold=True), box, Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[r].height = 34
            r += 1
            raw, nfs = b.get("raw"), b.get("nf")
            for i, row in enumerate(b["rows"]):
                for j, v in enumerate(row[:ncols]):
                    rv = raw[i][j] if raw and j < len(raw[i]) else None
                    c = _cell(ws, r, j + 1, rv if rv is not None else v)
                    c.font, c.border = F(), box
                    if rv is not None:
                        c.number_format = nfs[i] if nfs else "0.00"
                        c.alignment = Alignment(horizontal="right")
                    else:
                        c.alignment = Alignment(wrap_text=True, vertical="top")
                widths = [46] + [17] * 11
                lines = max([1] + [len(str(v)) // max(1, int(widths[j] * 0.95)) + 1 for j, v in enumerate(row[:ncols]) if (not raw or j >= len(raw[i]) or raw[i][j] is None)])
                if lines > 1:
                    ws.row_dimensions[r].height = 16 * lines
                r += 1
            r += 1
        elif t == "chart":
            text_row(b["fig"], F(bold=True))
            hdr = r
            _cell(ws, hdr, 1, "Kỳ").font = F(bold=True)
            ws.row_dimensions[hdr].height = 48
            for j, s in enumerate(b["series"]):
                hc = _cell(ws, hdr, 2 + j, s["name"])
                hc.font, hc.alignment = F(bold=True), Alignment(wrap_text=True, vertical="top")
            for i, lab in enumerate(b["labels"]):
                _cell(ws, hdr + 1 + i, 1, lab).font = F()
                for j, s in enumerate(b["series"]):
                    vc = _cell(ws, hdr + 1 + i, 2 + j, s["values"][i])
                    vc.font, vc.number_format = F(), "#,##0.00"
            last = hdr + len(b["labels"])
            ch = LineChart() if b["kind"] == "line" else BarChart()
            if b["kind"] == "stacked":
                ch.grouping, ch.overlap = "stacked", 100
            ch.title = b["title"]
            ch.height, ch.width = 8.5, 17
            ch.add_data(Reference(ws, min_col=2, min_row=hdr, max_col=1 + len(b["series"]), max_row=last), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=hdr + 1, max_row=last))
            ws.add_chart(ch, f"A{last + 2}")
            r = last + 2 + 18
    wb.properties.title = rep["title"]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


EXPORTERS = {"docx": export_docx, "pdf": export_pdf, "xlsx": export_xlsx}
