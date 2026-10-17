"""Đọc file CSV/Excel một cách an toàn: kiểm tra loại file, giới hạn kích thước, nhận diện dòng tiêu đề và sheet."""
import csv
import io
import zipfile
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import load_workbook

from ..config import get_settings
from .parsing import collapse_spaces, is_missing, preview_cell

EXCEL_EXT = {"xlsx", "xlsm"}
CSV_EXT = {"csv", "tsv", "txt"}


class LoadError(Exception):
    """Lỗi có thông báo tiếng Việt hiển thị trực tiếp cho người dùng."""


@dataclass
class Table:
    df: pd.DataFrame
    file_type: str
    sheet: Optional[str]
    row_numbers: List[int]  # số dòng trong file gốc, cùng thứ tự với df
    info: Dict[str, Any] = field(default_factory=dict)


# ------------------------------------------------------------------ nhận diện
def file_ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def detect_type(filename: str, raw: bytes) -> str:
    ext = file_ext(filename)
    if ext == "xls":
        raise LoadError("Định dạng .xls (Excel cũ) chưa được hỗ trợ. Hãy mở bằng Excel và lưu lại thành .xlsx.")
    if ext in EXCEL_EXT:
        if not raw.startswith(b"PK\x03\x04"):
            raise LoadError("Tệp có đuôi Excel nhưng nội dung không phải file .xlsx hợp lệ.")
        return "excel"
    if ext in CSV_EXT:
        if raw.startswith(b"PK\x03\x04"):
            raise LoadError("Tệp có đuôi CSV nhưng thực chất là file nén/Excel. Hãy đổi đúng đuôi .xlsx.")
        if b"\x00" in raw[:4096] and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            raise LoadError("Tệp CSV không phải văn bản (chứa dữ liệu nhị phân).")
        return "csv"
    raise LoadError("Chỉ hỗ trợ tệp .csv, .tsv, .txt, .xlsx, .xlsm.")


# ------------------------------------------------------------------ tiêu đề
def _nonempty(row: List[Any]) -> int:
    return sum(1 for v in row if not is_missing(v))


def find_header(rows: List[List[Any]]) -> Optional[int]:
    """Chọn dòng tiêu đề: dòng đầu tiên có số ô không rỗng gần bằng số cột lớn nhất trong 15 dòng đầu."""
    head = rows[:15]
    counts = [_nonempty(r) for r in head]
    mx = max(counts) if counts else 0
    if mx == 0:
        return None
    threshold = 1 if mx < 2 else max(2, int(0.6 * mx + 0.999))
    for i, c in enumerate(counts):
        if c >= threshold:
            return i
    return None


def _clean_header(values: List[Any]) -> List[str]:
    names: List[str] = []
    seen: Dict[str, int] = {}
    for i, v in enumerate(values):
        name = collapse_spaces(str(v)) if not is_missing(v) else f"Cột {i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        names.append(name)
    return names


# ------------------------------------------------------------------ CSV
def _decode(raw: bytes) -> (str, str):
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16"), "utf-16"
    for enc in ("utf-8-sig", "utf-8", "cp1258"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1"), "latin-1"


def _sniff_delimiter(text: str) -> str:
    lines = [ln for ln in text.splitlines()[:30] if ln.strip()]
    best, best_score = ",", (0, 0)
    for d in (",", ";", "\t", "|"):
        counts = [len(r) for r in csv.reader(lines, delimiter=d)]
        if not counts:
            continue
        mode = max(set(counts), key=counts.count)
        score = (mode if mode > 1 else 0, counts.count(mode))
        if score > best_score:
            best, best_score = d, score
    return best


def _read_csv_rows(raw: bytes) -> (List[List[Any]], Dict[str, Any]):
    text, enc = _decode(raw)
    delim = _sniff_delimiter(text)
    try:
        rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    except csv.Error as exc:
        raise LoadError(f"Không đọc được tệp CSV: {exc}")
    return rows, {"encoding": enc, "delimiter": {",": "dấu phẩy", ";": "dấu chấm phẩy", "\t": "tab", "|": "gạch đứng"}[delim]}


# ------------------------------------------------------------------ Excel
def _open_workbook(raw: bytes):
    try:
        zf = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile:
        raise LoadError("File Excel bị hỏng hoặc không đúng định dạng .xlsx.")
    if sum(i.file_size for i in zf.infolist()) > 200 * 1024 * 1024:
        raise LoadError("File Excel sau khi giải nén quá lớn (trên 200 MB).")
    try:
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise LoadError("Không đọc được file Excel (có thể bị hỏng hoặc đặt mật khẩu).")


def _sheet_summary(title: str, sample: List[List[Any]], max_row: Optional[int], state: str) -> Dict[str, Any]:
    hdr = find_header(sample)
    info: Dict[str, Any] = {"name": title, "hidden": state != "visible", "tabular": False, "n_rows": 0, "n_cols": 0, "note": ""}
    if hdr is None:
        info["note"] = "Sheet trống."
        return info
    header = sample[hdr]
    width = max((i + 1 for i, v in enumerate(header) if not is_missing(v)), default=0)
    data = sample[hdr + 1:]
    n_data = max(0, (max_row or (hdr + 1 + len(data))) - (hdr + 1))
    fills = [_nonempty(r[:width]) / width for r in data if width and _nonempty(r)]
    fill = sum(fills) / len(fills) if fills else 0
    info.update(n_rows=n_data, n_cols=width, header_row=hdr + 1)
    info["tabular"] = width >= 2 and n_data >= 3 and fill >= 0.3
    if not info["tabular"]:
        info["note"] = "Có vẻ không phải bảng dữ liệu (ít cột hoặc ít dòng), thường là ghi chú."
    info["header"] = [preview_cell(v) for v in header[:width]]
    return info


def inspect_file(raw: bytes, filename: str) -> Dict[str, Any]:
    ftype = detect_type(filename, raw)
    if ftype == "csv":
        rows, meta = _read_csv_rows(raw)
        s = _sheet_summary("CSV", rows[:40], len(rows), "visible")
        s.update(meta)
        return {"file_type": "csv", "sheets": [], "default_sheet": None, "csv": s}
    wb = _open_workbook(raw)
    sheets = []
    try:
        for ws in wb.worksheets:
            if not hasattr(ws, "iter_rows"):
                continue
            sample = [list(r) for r in ws.iter_rows(min_row=1, max_row=40, values_only=True)]
            max_row = ws.max_row if ws.max_row and ws.max_row < 2_000_000 else None
            sheets.append(_sheet_summary(ws.title, sample, max_row, ws.sheet_state))
    finally:
        wb.close()
    if not sheets:
        raise LoadError("File Excel không có sheet dữ liệu nào.")
    tab = [s for s in sheets if s["tabular"] and not s["hidden"]] or [s for s in sheets if s["tabular"]]
    default = max(tab, key=lambda s: s["n_rows"] * s["n_cols"])["name"] if tab else sheets[0]["name"]
    return {"file_type": "excel", "sheets": sheets, "default_sheet": default, "csv": None}


# ------------------------------------------------------------------ đọc bảng
def read_table(raw: bytes, filename: str, sheet: Optional[str] = None, limit: Optional[int] = None) -> Table:
    settings = get_settings()
    ftype = detect_type(filename, raw)
    info: Dict[str, Any] = {}
    all_rows: List[List[Any]] = []

    if ftype == "csv":
        all_rows, meta = _read_csv_rows(raw)
        info.update(meta)
        used_sheet = None
        first_row_no = 1
        if limit:
            all_rows = all_rows[: limit + 20]
    else:
        wb = _open_workbook(raw)
        try:
            names = [w.title for w in wb.worksheets if hasattr(w, "iter_rows")]
            if sheet is None:
                sheet = inspect_file(raw, filename)["default_sheet"]
            if sheet not in names:
                raise LoadError(f"Không tìm thấy sheet '{sheet}' trong file.")
            ws = wb[sheet]
            cap = (limit + 20) if limit else settings.max_rows + 25
            for r in ws.iter_rows(values_only=True):
                all_rows.append(list(r))
                if len(all_rows) > cap:
                    break
        finally:
            wb.close()
        used_sheet = sheet
        first_row_no = 1

    hdr = find_header(all_rows)
    if hdr is None:
        raise LoadError("Không tìm thấy dòng tiêu đề/dữ liệu trong tệp (tệp trống?).")
    header_vals = all_rows[hdr]
    data_rows = all_rows[hdr + 1:]
    row_numbers = [first_row_no + hdr + 1 + i for i in range(len(data_rows))]

    if limit:
        data_rows, row_numbers = data_rows[:limit], row_numbers[:limit]
    elif len(data_rows) > settings.max_rows:
        raise LoadError(
            f"Tệp có hơn {settings.max_rows:,} dòng, vượt giới hạn của dịch vụ miễn phí. Hãy chia nhỏ tệp rồi tải lại."
        )

    width = max((i + 1 for i, v in enumerate(header_vals) if not is_missing(v)), default=0)
    for r in data_rows:
        for i in range(len(r) - 1, width - 1, -1):
            if not is_missing(r[i]):
                width = max(width, i + 1)
                break
    if width == 0:
        raise LoadError("Không tìm thấy cột dữ liệu nào.")
    if width > settings.max_cols:
        raise LoadError(f"Tệp có {width} cột, vượt giới hạn {settings.max_cols} cột.")

    names = _clean_header((header_vals + [None] * width)[:width])
    fixed, fixed_nums, dropped_rows = [], [], 0
    for r, no in zip(data_rows, row_numbers):
        r = (list(r) + [None] * width)[:width]
        if _nonempty(r) == 0:
            dropped_rows += 1
            continue
        fixed.append(r)
        fixed_nums.append(no)
    if not fixed:
        raise LoadError("Tệp chỉ có dòng tiêu đề, không có dữ liệu.")

    df = pd.DataFrame(fixed, columns=names, dtype=object)
    keep_cols = [c for c in names if df[c].map(lambda v: not is_missing(v)).any() or not c.startswith("Cột ")]
    dropped_cols = len(names) - len(keep_cols)
    df = df[keep_cols]
    info.update(dropped_empty_rows=dropped_rows, dropped_empty_cols=dropped_cols, header_row=hdr + 1)
    return Table(df=df, file_type=ftype, sheet=used_sheet, row_numbers=fixed_nums, info=info)
