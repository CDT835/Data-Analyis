"""Báo cáo do người dùng tự soạn: chuẩn hóa khối, từ chối dữ liệu xấu, đánh số, xuất ba định dạng, chèn từ dữ liệu."""
import base64
import io

import pytest

from tests.test_datasets import call
from tests.test_finance import CLEAN, parse
from tests.test_reports import body, texts


@pytest.fixture(scope="module")
def clean():
    return parse(CLEAN)


DOC = [
    {"t": "heading", "level": 1, "text": "Mở đầu"},
    {"t": "para", "text": "Đoạn một.\nĐoạn hai có <b>thẻ</b> và \x07 ký tự lạ."},
    {"t": "heading", "level": 2, "text": "Bối cảnh"},
    {"t": "bullets", "items": ["Ý một", "", "Ý hai"]},
    {"t": "table", "headers": ["Chỉ tiêu", "2024", "2025"], "rows": [["Doanh thu", "100", "120"], ["Lợi nhuận", "10"]], "caption": "Bảng 9. Số liệu nhập tay"},
    {"t": "callout", "kind": "xyz", "text": "Lưu ý quan trọng"},
    {"t": "chart", "kind": "bar", "title": "Doanh thu", "labels": ["2024", "2025"], "series": [{"name": "DT", "values": [100, 120]}], "unit": "triệu"},
    {"t": "heading", "level": 1, "text": "Kết luận"},
]


def render(blocks, fmt="docx", **extra):
    return call("/api/v1/reports/render", {"title": "Báo cáo tự soạn", "org": "Đơn vị A", "blocks": blocks, "format": fmt, **extra})


def test_render_three_formats_and_numbering():
    for fmt in ("docx", "pdf", "xlsx"):
        r = render(DOC, fmt)
        assert r.status_code == 200, r.text
        d = r.json()
        raw = base64.b64decode(d["content_b64"])
        assert d["size"] == len(raw) > 1000 and d["filename"] == f"Báo_cáo_tự_soạn.{fmt}"
        assert (raw[:4] == b"%PDF") if fmt == "pdf" else raw[:2] == b"PK"
    import docx
    doc = docx.Document(io.BytesIO(base64.b64decode(render(DOC, "docx").json()["content_b64"])))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "1. Mở đầu" in text and "1.1. Bối cảnh" in text and "2. Kết luận" in text  # tự đánh số mục
    assert "Đoạn một." in text and "Đoạn hai có <b>thẻ</b> và  ký tự lạ." in text  # xuống dòng thành đoạn mới, bỏ ký tự điều khiển
    assert "Bảng 1. Số liệu nhập tay" in text and "Hình 1. Doanh thu" in text and "Bảng 9" not in text  # đánh số lại bảng, hình
    assert "Lưu ý: Lưu ý quan trọng" in text  # kiểu ghi chú lạ về "Lưu ý"
    t = doc.tables[0]
    assert [c.text for c in t.rows[2].cells] == ["Lợi nhuận", "10", ""]  # hàng thiếu ô được bổ sung


def test_excel_has_real_chart_and_pdf_has_text():
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(render(DOC, "xlsx").json()["content_b64"])))
    assert len(wb.active._charts) == 1
    from pypdf import PdfReader
    pdf = PdfReader(io.BytesIO(base64.b64decode(render(DOC, "pdf").json()["content_b64"])))
    assert "Báo cáo tự soạn".upper() in "".join(pg.extract_text() for pg in pdf.pages)


@pytest.mark.parametrize("blocks,expect", [
    ([], "chưa có nội dung"),
    ([{"t": "para", "text": "   "}, {"t": "heading", "text": ""}], "chưa có nội dung"),
    ([{"t": "script", "text": "x"}], "không được hỗ trợ"),
    ([{"t": "para", "text": "x"}] * 401, "quá nhiều mục"),
    ([{"t": "table", "headers": ["a"] * 13, "rows": []}], "Bảng quá lớn"),
    ([{"t": "table", "headers": ["a"], "rows": [["x"]] * 101}], "Bảng quá lớn"),
    ([{"t": "chart", "kind": "bar", "title": "t", "labels": ["a"], "series": [{"name": "s", "values": [1]}]}] * 11, "quá nhiều biểu đồ"),
    ([{"t": "chart", "kind": "bar", "title": "t", "labels": ["a"] * 41, "series": [{"name": "s", "values": [1]}]}], "quá nhiều điểm"),
])
def test_rejects_bad_documents(blocks, expect):
    r = render(blocks)
    assert r.status_code == 422 and expect in r.json()["detail"], r.text


def test_invalid_shapes_are_dropped_not_crashing():
    blocks = [{"t": "para", "text": "Còn lại"}, {"t": "table", "headers": [], "rows": []}, {"t": "chart", "kind": "pie", "labels": ["a"], "series": []},
              {"t": "table", "headers": ["a", "b"], "rows": [["1", "2"]], "raw": [[1, "x"]], "nf": ["0"], "num_cols": [5]},
              {"t": "chart", "kind": "line", "title": "", "labels": ["a", "b"], "series": [{"name": "s", "values": [1, float("nan") if False else "bad", 3]}]}]
    r = render(blocks)
    assert r.status_code == 200, r.text
    assert call("/api/v1/reports/render", {"title": "", "blocks": blocks, "format": "docx"}).status_code == 200
    assert render(blocks + ["rác"], "docx").status_code == 422  # phần tử không phải khối
    assert render(blocks, "exe").status_code == 422  # định dạng lạ
    assert call("/api/v1/reports/render", {"blocks": "x", "format": "pdf"}).status_code == 422


def test_numeric_cells_keep_excel_numbers():
    import openpyxl
    blk = [{"t": "table", "headers": ["Kỳ", "Giá trị"], "rows": [["2025", "1.234"]], "raw": [[None, 1234.5]], "nf": ["#,##0.0"], "caption": "Số"}]
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(render(blk, "xlsx").json()["content_b64"])))
    cells = {c.value for row in wb.active.iter_rows() for c in row}
    assert 1234.5 in cells


def test_requires_signature():
    import httpx  # noqa: F401
    from fastapi.testclient import TestClient
    from app.main import app
    assert TestClient(app).post("/api/v1/reports/render", json={"blocks": DOC, "format": "pdf"}).status_code in (401, 403)


def test_fragment_mode_returns_unnumbered_blocks_for_insertion(clean):
    r = call("/api/v1/reports/preview", body(clean, fragment=True, sections=["kpis", "liquidity", "revenue_profit_chart"]))
    assert r.status_code == 200, r.text
    rep = r.json()
    kinds = [b["t"] for b in rep["blocks"]]
    assert "title" not in kinds and rep["included"] == ["kpis", "revenue_profit_chart", "liquidity"]  # theo thứ tự trong mẫu
    assert all("num" not in b for b in rep["blocks"]) and all(not (b.get("caption") or "").startswith("Bảng") for b in rep["blocks"])
    assert any(b["t"] == "chart" and "fig" not in b for b in rep["blocks"])
    # khối chèn vào rồi xuất lại được và được đánh số
    r2 = render(rep["blocks"], "docx")
    assert r2.status_code == 200, r2.text
    # chế độ thường vẫn như cũ
    full = call("/api/v1/reports/preview", body(clean, sections=["kpis", "liquidity"])).json()
    assert full["blocks"][0]["t"] == "title" and any(b["t"] == "heading" and b["num"] == "1" for b in full["blocks"])
    assert any((b.get("caption") or "").startswith("Bảng 1.") for b in full["blocks"])


def test_inserted_data_numbers_match_analysis(clean):
    """Số trong khối chèn từ dữ liệu giống hệt khối của báo cáo dựng tự động (không có con số mới)."""
    frag = call("/api/v1/reports/preview", body(clean, fragment=True, sections=["liquidity"])).json()
    full = call("/api/v1/reports/preview", body(clean, sections=["liquidity"])).json()
    tf = [b["rows"] for b in frag["blocks"] if b["t"] == "table"]
    tu = [b["rows"] for b in full["blocks"] if b["t"] == "table"]
    assert tf == tu and tf


def test_text_columns_of_hand_made_tables_are_left_aligned():
    from app.services.report_custom import normalize_document
    doc = normalize_document("T", "", None, [{"t": "table", "headers": ["Hạng mục", "Kế hoạch", "Ghi chú"], "rows": [["Giảm tồn kho", "10%", "Quý 1"], ["Thu nợ", "15%", "Quý 2"]]}])
    tb = [b for b in doc["blocks"] if b["t"] == "table"][0]
    assert tb["num_cols"] == [1]  # chỉ cột toàn số mới căn phải
    doc = normalize_document("T", "", None, [{"t": "table", "headers": ["A", "B"], "rows": [["x", "y"]]}])
    assert [b for b in doc["blocks"] if b["t"] == "table"][0]["num_cols"] == []
