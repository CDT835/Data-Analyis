import base64
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services import finance
from tests.test_datasets import DATA, STUDENTS, call, payload

CLEAN = DATA / "bao_cao_tai_chinh_mau_2021_2025.xlsx"
FLAWED = DATA / "bao_cao_tai_chinh_co_loi.xlsx"
CSV_BS = DATA / "bang_can_doi_ke_toan.csv"


def parse(p: Path):
    r = call("/api/v1/finance/parse", payload(p))
    assert r.status_code == 200, r.text
    return r.json()


def sheet_values(p: Path, sheet: str):
    """Đọc độc lập bằng openpyxl: mã số -> 5 giá trị."""
    ws = load_workbook(p, data_only=True)[sheet]
    out = {}
    for row in ws.iter_rows(min_row=6, values_only=True):
        if row[1] is not None:
            out[str(int(row[1]))] = [row[3 + i] for i in range(5)]
    return out


def ratio(res, key):
    for g in res["analysis"]["ratio_groups"]:
        for it in g["items"]:
            if it["key"] == key:
                return it
    raise KeyError(key)


def test_clean_file_structure_and_mapping():
    r = parse(CLEAN)
    sts = {s["type"]: s for s in r["parsed"]["statements"]}
    assert set(sts) == {"bs", "is"} and sts["bs"]["unit"] == "triệu đồng"
    assert [p["label"] for p in sts["bs"]["periods"]] == ["2021", "2022", "2023", "2024", "2025"]
    assert [s["sheet"] for s in r["parsed"]["skipped"]] == ["Ghi chú"]  # sheet ghi chú không bị nhầm là báo cáo
    m = r["mapping"]
    assert m["total_assets"]["code"] == "270" and m["net_profit"]["code"] == "60" and m["interest_expense"]["code"] == "23"
    assert r["unmapped_important"] == {}


def test_clean_file_balances_and_all_checks_pass():
    c = parse(CLEAN)["checks"]
    assert c["n_failed"] == 0 and c["n_checked"] >= 60 and not c["data_issues"]


def test_ratios_match_independent_calculation():
    r = parse(CLEAN)
    bs, is_ = sheet_values(CLEAN, "Bảng cân đối kế toán"), sheet_values(CLEAN, "Kết quả kinh doanh")
    cr = ratio(r, "current_ratio")["values"]
    de = ratio(r, "debt_to_equity")["values"]
    gm = ratio(r, "gross_margin")["values"]
    roe = ratio(r, "roe")
    dso = ratio(r, "dso")["values"]
    for i in range(5):
        assert cr[i] == pytest.approx(bs["100"][i] / bs["310"][i], abs=1e-3)
        assert de[i] == pytest.approx(bs["300"][i] / bs["400"][i], abs=1e-3)
        assert gm[i] == pytest.approx(is_["20"][i] / is_["10"][i] * 100, abs=1e-2)
    # ROE dùng vốn chủ sở hữu bình quân; kỳ đầu dùng số cuối kỳ
    assert roe["values"][0] == pytest.approx(is_["60"][0] / bs["400"][0] * 100, abs=1e-2) and roe["basis"][0] == "cuối kỳ"
    for i in range(1, 5):
        assert roe["values"][i] == pytest.approx(is_["60"][i] / ((bs["400"][i] + bs["400"][i - 1]) / 2) * 100, abs=1e-2)
        assert roe["basis"][i] == "bình quân"
        assert dso[i] == pytest.approx(365 / (is_["10"][i] / ((bs["130"][i] + bs["130"][i - 1]) / 2)), abs=1e-2)


def test_dupont_reproduces_roe():
    r = parse(CLEAN)["analysis"]
    roe = next(it for g in r["ratio_groups"] for it in g["items"] if it["key"] == "roe")["values"]
    for row, v in zip(r["dupont"]["rows"], roe):
        assert row["roe_pct"] == pytest.approx(v, abs=0.01)


def test_alerts_are_grounded_in_the_numbers():
    a = parse(CLEAN)["analysis"]["alerts"]
    titles = {(x["period"], x["title"]) for x in a}
    assert ("2023", "Thanh toán ngắn hạn thấp") in titles and ("2024", "Thanh toán ngắn hạn thấp") in titles
    assert ("2024", "Biên lợi nhuận gộp giảm") in titles and ("2024", "Lợi nhuận sau thuế giảm mạnh") in titles
    assert not any(x["period"] in ("2021", "2022") for x in a)  # kỳ khỏe mạnh không bị cảnh báo
    assert not any(x["severity"] == "high" for x in a)
    assert all("ngưỡng tham khảo" in x["message"] or "%" in x["message"] or "ngày" in x["message"] or "lần" in x["message"] for x in a)


def test_flawed_file_reports_exact_root_causes():
    c = parse(FLAWED)["checks"]
    roots = {(x["title"], x["period"]): x["diff"] for x in c["failed"] if not x["cascade"]}
    assert roots == {("Tổng cộng tài sản", "2023"): 500.0, ("Tài sản ngắn hạn", "2024"): 300.0,
                     ("Lợi nhuận gộp", "2022"): -200.0, ("Lợi nhuận sau thuế", "2025"): 150.0}
    assert c["n_failed"] == 7 and c["n_root_failed"] == 4  # 3 lỗi còn lại là hệ quả của lỗi gốc
    assert [x["cascade"] for x in c["failed"]] == [False] * 4 + [True] * 3  # lỗi gốc xếp trước
    msgs = " ".join(i["message"] for i in c["data_issues"])
    assert "1 ô trống" in msgs and "Hàng tồn kho (2021)" in msgs and "2 ô số đang lưu dạng chữ" in msgs


def test_text_numbers_with_parentheses_are_read():
    assert finance.fin_number("(22.000)") == (-22000.0, "text")
    assert finance.fin_number("1.234.567") == (1234567.0, "text")
    assert finance.fin_number("-") == (0.0, "dash") and finance.fin_number(None) == (None, None)
    assert finance.fin_number("abc") == (None, "invalid")
    bs = sheet_values(FLAWED, "Bảng cân đối kế toán")
    r = parse(FLAWED)
    # ô lưu dạng chữ "(xx.xxx)" vẫn được đọc đúng thành số âm
    row = next(x for x in r["parsed"]["statements"][0]["rows"] if x["code"] == "223")
    assert row["values"][2] < 0 and row["values"][2] == pytest.approx(sheet_values(CLEAN, "Bảng cân đối kế toán")["223"][2])


def test_csv_with_only_balance_sheet():
    r = parse(CSV_BS)
    assert [s["type"] for s in r["parsed"]["statements"]] == ["bs"]
    assert set(r["unmapped_important"]) >= {"net_revenue", "cogs", "net_profit"}
    assert r["checks"]["n_failed"] == 0 and not r["checks"]["data_issues"]
    npm = ratio(r, "net_margin")
    assert not npm["available"] and "Doanh thu thuần" in npm["reason"] and ratio(r, "current_ratio")["available"]


def test_periods_and_units():
    assert finance.period_info("31/12/2025")["pid"] == "2025" and finance.period_info("Năm 2024")["pid"] == "2024"
    assert finance.period_info("Quý 2/2025")["pid"] == "2025-Q2" and finance.period_info("Tháng 3/2025")["pid"] == "2025-03"
    assert finance.period_info("Số cuối năm")["kind"] == "end" and finance.period_info("Số đầu năm")["kind"] == "start"
    assert finance.period_info("BÁO CÁO KẾT QUẢ HOẠT ĐỘNG KINH DOANH NĂM 2025") is None  # tiêu đề dài không bị nhầm thành cột kỳ
    assert finance.period_info("Ghi chú") is None and finance.period_info(2025)["pid"] == "2025"
    assert finance.name_key("   - Nguyên giá (*)") == "nguyen gia" and finance.name_key("A - TÀI SẢN NGẮN HẠN") == "tai san ngan han"


def test_negative_costs_are_normalised():
    csv_text = "Chỉ tiêu,Mã số,2024,2025\nDoanh thu thuần,10,100.000,120.000\nGiá vốn hàng bán,11,(70.000),(80.000)\nLợi nhuận gộp,20,30.000,40.000\n"
    p = {"filename": "kq.csv", "file_b64": base64.b64encode(csv_text.encode("utf-8")).decode()}
    r = call("/api/v1/finance/parse", p).json()
    assert any("đổi dấu" in n for n in r["analysis"]["notes"])
    assert r["checks"]["n_failed"] == 0
    assert ratio(r, "gross_margin")["values"] == [pytest.approx(30.0), pytest.approx(33.3333, abs=1e-3)]


def test_override_mapping_recalculates_without_file():
    r = parse(CLEAN)
    bs = next(s for s in r["parsed"]["statements"] if s["type"] == "bs")
    idx = next(i for i, x in enumerate(bs["rows"]) if x["code"] == "141")
    r2 = call("/api/v1/finance/analyze", {"parsed": r["parsed"], "overrides": {"inventory": {"statement": bs["id"], "row": idx}}}).json()
    assert r2["mapping"]["inventory"]["source"] == "bạn chọn" and r2["mapping"]["inventory"]["code"] == "141"
    assert ratio(r2, "quick_ratio")["values"][0] != ratio(r, "quick_ratio")["values"][0]
    r3 = call("/api/v1/finance/analyze", {"parsed": r["parsed"], "overrides": {"inventory": None}}).json()
    assert "inventory" not in r3["mapping"] and not ratio(r3, "quick_ratio")["available"]
    bad = call("/api/v1/finance/analyze", {"parsed": {"statements": "x"}, "overrides": {}})
    assert bad.status_code == 422


def test_rejects_non_financial_files():
    r = call("/api/v1/finance/parse", payload(STUDENTS))
    assert r.status_code == 422 and "Không nhận diện được" in r.json()["detail"]
    assert call("/api/v1/finance/parse", payload(DATA / "ban_hang_2025_2026.xlsx")).status_code == 422
