import base64
import io
import json

import pytest

from app.services.ai import vn, vnum
from tests.test_ai import faithful, prompt_of
from tests.test_datasets import call
from tests.test_finance import CLEAN, FLAWED, parse


@pytest.fixture(scope="module")
def clean():
    return parse(CLEAN)


@pytest.fixture(scope="module")
def flawed():
    return parse(FLAWED)


def body(r, **extra):
    stm = [{"type": s["type"], "type_label": s["type_label"], "sheet": s["sheet"], "unit": s["unit"], "n_periods": len(s["periods"]), "n_rows": len(s["rows"])} for s in r["parsed"]["statements"]]
    d = {"template": "fin_overview", "title": "Báo cáo thử nghiệm", "org": "Công ty An Phát", "analysis": r["analysis"], "checks": r["checks"],
         "mapping": r["mapping"], "std_items": r["std_items"], "statements": stm}
    d.update(extra)
    return d


def preview(r, **extra):
    resp = call("/api/v1/reports/preview", body(r, **extra))
    return resp


def texts(rep):
    out = []
    for b in rep["blocks"]:
        for k in ("text", "caption", "fig", "subtitle"):
            if b.get(k):
                out.append(b[k])
        out += b.get("items", []) if b["t"] == "bullets" else []
        for row in b.get("rows", []):
            out += [str(x) for x in row]
        if b["t"] == "kpis":
            out += [f"{i['label']} {i['value']} {i['delta']}" for i in b["items"]]
    return "\n".join(out)


def fake_ai(clean_, extra_text="", edited=None, stale=False):
    facts = call("/api/v1/ai/facts", {"analysis": clean_["analysis"], "checks": clean_["checks"]}).json()["facts"]
    prompt = "\n".join(f"{f['id']}: {f['text']}" for f in facts)
    from app.services import ai
    com, ver = ai.validate(faithful(prompt, extra_text), facts)
    return {"commentary": com, "facts": facts, "model": "openai/gpt-oss-120b", "disclaimer": "Nhận xét do AI tạo.", "edited": edited, "stale": stale}, facts


def test_templates_and_availability(clean):
    d = call("/api/v1/reports/templates", {"mapping": clean["mapping"], "statements": [{"type": "bs"}, {"type": "is"}]}).json()
    assert [t["slug"] for t in d["templates"]] == ["fin_overview", "fin_liquidity", "fin_profitability", "fin_quality"]
    assert all(t["usable"] for t in d["templates"]) and sum(len(t["sections"]) for t in d["templates"] if t["slug"] == "fin_overview") == 18
    only_bs = call("/api/v1/reports/templates", {"statements": [{"type": "bs"}]}).json()
    byslug = {t["slug"]: t for t in only_bs["templates"]}
    assert byslug["fin_liquidity"]["usable"] and not byslug["fin_profitability"]["usable"] and "Báo cáo kết quả kinh doanh" in byslug["fin_profitability"]["reason"]
    assert not next(s for s in byslug["fin_overview"]["sections"] if s["slug"] == "revenue_profit_chart")["usable"]


def test_blocks_match_computed_numbers(clean):
    rep = preview(clean).json()
    a = clean["analysis"]
    txt = texts(rep)
    latest = a["latest"]
    assert vnum(latest["net_revenue"]["value"]) in txt and vnum(latest["net_profit"]["value"]) in txt
    cur = next(i for g in a["ratio_groups"] for i in g["items"] if i["key"] == "current_ratio")
    row = next(r for b in rep["blocks"] if b["t"] == "table" for r in b["rows"] if r[0].startswith("Tỷ số thanh toán hiện hành"))
    assert row[1:] == [vn(v, 2) for v in cur["values"]]
    for dp in a["dupont"]["rows"]:
        assert vn(dp["roe_pct"], 1) + "%" in txt
    heads = [b for b in rep["blocks"] if b["t"] == "heading" and b["level"] == 1]
    assert [h["num"] for h in heads] == [str(i + 1) for i in range(len(heads))]
    assert rep["blocks"][0]["t"] == "title" and rep["title"] == "Báo cáo thử nghiệm"
    assert 0 < rep["n_charts"] <= 10 and sum(1 for b in rep["blocks"] if b["t"] == "chart") == rep["n_charts"]
    capital = next(b for b in rep["blocks"] if b["t"] == "chart" and b["kind"] == "stacked")
    for i in range(len(capital["labels"])):
        assert abs(sum(s["values"][i] for s in capital["series"]) - 100) < 0.5  # cột chồng cộng đủ 100%
    figs = [b["fig"] for b in rep["blocks"] if b["t"] == "chart"]
    assert figs[0].startswith("Hình 1.")


@pytest.mark.parametrize("tpl", ["fin_overview", "fin_liquidity", "fin_profitability", "fin_quality"])
def test_every_template_builds(clean, tpl):
    r = preview(clean, template=tpl)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["included"] and d["template"] == tpl


def test_ai_sections_skipped_without_comment_then_included(clean):
    d = preview(clean).json()
    sk = {s["section"]: s["reason"] for s in d["skipped"]}
    assert {"ai_summary", "ai_detail", "ai_reco"} <= set(sk) and "AI nhận xét" in sk["ai_summary"]
    assert "Các phần đã bỏ qua" in texts(d)
    ai, _ = fake_ai(clean)
    d2 = preview(clean, ai=ai).json()
    assert not d2["skipped"] and "Dẫn chứng: F" in texts(d2) and "Tóm tắt của AI" in texts(d2)


def test_edited_comment_is_shown_and_reverified(clean):
    ai, _ = fake_ai(clean)
    edited = {"summary": {"text": "Doanh thu tăng 97,3% nhờ chiến dịch lớn.", "refs": []}, "highlights": [{"text": "Biên lợi nhuận được cải thiện.", "refs": []}],
              "trends": [], "risks": [], "recommendations": [], "next_analyses": []}
    ai["edited"] = edited
    d = preview(clean, ai=ai).json()
    t = texts(d)
    assert "97,3%" in t and "chưa đối chiếu được" in t and "đã được người dùng chỉnh sửa" in t
    assert ai["commentary"]["summary"]["text"] not in t  # hiển thị bản đã sửa, không phải bản gốc
    stale = preview(clean, ai=fake_ai(clean, stale=True)[0]).json()
    assert "phiên bản phân tích cũ hơn" in texts(stale)
    flagged, _ = fake_ai(clean, extra_text=" Tăng 987,6% bất thường.")
    assert "chưa đối chiếu được" in texts(preview(clean, ai=flagged).json())


def test_unusable_or_invalid_requests_are_rejected(clean):
    only_bs = body(clean, template="fin_profitability", statements=[{"type": "bs"}], mapping={k: v for k, v in clean["mapping"].items() if v["statement"].startswith("bs:")})
    r = call("/api/v1/reports/preview", only_bs)
    assert r.status_code == 422 and "Báo cáo kết quả kinh doanh" in r.json()["detail"]
    assert preview(clean, template="khong_co").status_code == 422
    assert preview(clean, sections=[]).status_code == 422
    assert preview(clean, template="fin_quality", sections=["liquidity"]).status_code == 422  # phần không thuộc mẫu
    assert preview(clean, title="x" * 201).status_code == 422
    assert call("/api/v1/reports/preview", {"template": "fin_overview", "analysis": {"bad": 1}}).status_code == 422
    assert call("/api/v1/reports/export", body(clean, format="exe")).status_code == 422
    assert call("/api/v1/reports/preview", body(clean), sign=False).status_code == 401


def test_missing_data_sections_are_skipped_with_reason(clean):
    bs_map = {k: v for k, v in clean["mapping"].items() if v["statement"].startswith("bs:")}
    d = preview(clean, statements=[{"type": "bs"}], mapping=bs_map).json()
    sk = {s["section"]: s["reason"] for s in d["skipped"]}
    assert "Báo cáo kết quả kinh doanh" in sk["revenue_profit_chart"] and "Báo cáo kết quả kinh doanh" in sk["profitability"]
    assert "liquidity" in d["included"] and "Biểu đồ doanh thu" not in texts(d).split("Các phần đã bỏ qua")[0]


def test_flawed_data_quality_template(flawed):
    d = preview(flawed, template="fin_quality").json()
    t = texts(d)
    assert "7 phép bị lệch" in t and "4 lỗi gốc" in t and "3 hệ quả" in t and "Số liệu có sai lệch" in t
    tbl = next(b for b in d["blocks"] if b["t"] == "table" and b.get("caption", "").endswith("bị lệch"))
    assert sum(1 for r in tbl["rows"] if r[3] == "Lỗi gốc") == 4 and sum(1 for r in tbl["rows"] if r[3] == "Hệ quả") == 3
    assert "Có 1 ô trống" in t  # vấn đề dữ liệu đầu vào


def test_title_is_sanitised(clean):
    d = preview(clean, title="Báo\x00 cáo\x07 mẫu", org="  Đơn vị\x1b A ").json()
    assert d["title"] == "Báo cáo mẫu" and d["org"] == "Đơn vị A"


# ----------------------------------------------------------------------------- xuất tệp
def export(clean, fmt, **extra):
    r = call("/api/v1/reports/export", body(clean, format=fmt, **extra))
    assert r.status_code == 200, r.text
    d = r.json()
    raw = base64.b64decode(d["content_b64"])
    assert d["size"] == len(raw) and d["filename"].endswith("." + fmt)
    return d, raw


def test_export_docx_roundtrip(clean):
    from docx import Document
    ai, _ = fake_ai(clean)
    d, raw = export(clean, "docx", ai=ai)
    assert raw[:2] == b"PK" and d["filename"] == "Báo_cáo_thử_nghiệm.docx".replace("á", "á")
    doc = Document(io.BytesIO(raw))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "BÁO CÁO THỬ NGHIỆM" in text and "1. Thông tin báo cáo" in text and "Hình 1." in text
    cells = " ".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert vnum(clean["analysis"]["latest"]["net_revenue"]["value"]) in cells
    assert len(doc.inline_shapes) >= 4 and len(doc.tables) >= 10
    assert doc.styles["Normal"].font.name == "Times New Roman" and doc.styles["Heading 1"].font.name == "Times New Roman"
    first = next(t for t in doc.tables if len(t.columns) >= 5 and t.rows[0].cells[0].text == 'Chỉ số')
    assert first.columns[0].width > first.columns[1].width  # cột đầu rộng hơn các cột kỳ


def test_export_pdf_roundtrip(clean):
    from pypdf import PdfReader
    d, raw = export(clean, "pdf")
    assert raw[:5] == b"%PDF-"
    r = PdfReader(io.BytesIO(raw))
    text = "\n".join(p.extract_text() for p in r.pages)
    assert len(r.pages) >= 4 and "Khả năng thanh toán" in text and "Tỷ số thanh toán hiện hành" in text and "137.120" in text


def test_export_xlsx_roundtrip(clean):
    from openpyxl import load_workbook
    d, raw = export(clean, "xlsx")
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Báo cáo"]
    nums = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float))]
    assert 137120 in nums and len(ws._charts) >= 4  # số ghi dạng số, biểu đồ Excel thật
    assert any(isinstance(c.value, str) and "Khả năng thanh toán" in c.value for row in ws.iter_rows() for c in row)
    assert ws.cell(1, 1).font.name == "Times New Roman"


def test_export_reports_skipped_parts_and_rejects_bad_template(clean, flawed):
    d, _ = export(clean, "docx")
    assert {s["section"] for s in d["skipped"]} == {"ai_summary", "ai_detail", "ai_reco"}
    d2, _ = export(flawed, "pdf", template="fin_quality")
    assert d2["included"] == ["info", "scope", "data_checks", "mapping", "alerts", "notes"]
    assert call("/api/v1/reports/export", body(clean, format="pdf", template="nope")).status_code == 422
