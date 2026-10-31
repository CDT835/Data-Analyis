"""Bảo mật dữ liệu vào/ra: chống tiêm công thức Excel, từ chối tệp macro và tệp 'bom nén'."""
import base64
import io
import zipfile

import openpyxl
import pytest

from app.services import loader
from tests.test_datasets import call
from tests.test_reports_custom import render

EVIL = ['=HYPERLINK("http://evil.test","bấm")', "+1+1", "-2+3", "@SUM(1+1)", "\t=1+1", "=cmd|' /C calc'!A0"]


def test_excel_export_never_creates_formulas():
    blocks = [
        {"t": "heading", "level": 1, "text": EVIL[0]},
        {"t": "para", "text": EVIL[1]},
        {"t": "bullets", "items": [EVIL[2], EVIL[3]]},
        {"t": "table", "headers": [EVIL[0], "B"], "rows": [[EVIL[4], EVIL[5]], ["Bình thường", "123"]]},
    ]
    r = render(blocks, "xlsx", title=EVIL[0], org=EVIL[3])
    assert r.status_code == 200, r.text
    raw = base64.b64decode(r.json()["content_b64"])
    # 1) không có phần tử <f> (công thức) nào trong sheet
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
        assert sheets and all(b"<f>" not in z.read(n) and b"<f " not in z.read(n) for n in sheets)
    # 2) mọi ô chuỗi nguy hiểm đều là văn bản và giữ nguyên nội dung gốc
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    seen = []
    for row in wb.active.iter_rows():
        for c in row:
            assert c.data_type != "f", f"{c.coordinate} bị coi là công thức: {c.value!r}"
            if isinstance(c.value, str) and c.value[:1] in "=+-@\t":
                seen.append(c.value)
                assert c.data_type == "s"
    assert any(v.startswith("=HYPERLINK") for v in seen) and "+1+1" in seen and any(v.startswith("@SUM") for v in seen)
    assert any(c.value == "Bình thường" for row in wb.active.iter_rows() for c in row)  # văn bản thường không bị đổi


def _xlsx_with(extra):
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    wb.active.append([1, 2])
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n in src.namelist():
            z.writestr(n, src.read(n))
        for name, data in extra.items():
            z.writestr(name, data)
    return out.getvalue()


def test_valid_xlsx_still_opens():
    assert loader._open_workbook(_xlsx_with({})) is not None


def test_xlsm_and_macro_extensions_rejected():
    for ext in ("xlsm", "xlsb", "xltm"):
        with pytest.raises(loader.LoadError, match="macro"):
            loader.detect_type(f"bao_cao.{ext}", b"PK\x03\x04")
    assert loader.detect_type("ok.xlsx", b"PK\x03\x04") == "excel"
    with pytest.raises(loader.LoadError):
        loader.detect_type("x.exe", b"MZ")


def test_xlsx_containing_vba_project_rejected():
    with pytest.raises(loader.LoadError, match="macro"):
        loader._open_workbook(_xlsx_with({"xl/vbaProject.bin": b"\x00" * 100}))


def test_zip_bomb_rejected():
    bomb = _xlsx_with({"xl/worksheets/sheet9.xml": b"0" * (8 * 1024 * 1024)})  # 8 MB số 0 nén còn vài KB
    with pytest.raises(loader.LoadError, match="bom nén"):
        loader._open_workbook(bomb)


def test_too_many_entries_rejected():
    many = _xlsx_with({f"xl/x/{i}.xml": b"x" for i in range(2100)})
    with pytest.raises(loader.LoadError, match="thành phần"):
        loader._open_workbook(many)


def test_upload_xlsm_rejected_by_api():
    payload = {"filename": "co_macro.xlsm", "file_b64": base64.b64encode(_xlsx_with({})).decode()}
    r = call("/api/v1/datasets/inspect", payload)
    assert r.status_code in (400, 422), r.text
    assert "macro" in r.text
